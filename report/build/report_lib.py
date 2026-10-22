"""Shared helpers and data for the OrbitRisk IEEE report builder."""
import sys
from pathlib import Path

import docx
import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
REPO = Path(r"C:\Users\Navsssssss\Desktop\orbitrisk")
PROC = REPO / "data" / "processed"


d = docx.Document(HERE / "template.docx")
P = d.paragraphs

# ---------------- data -----------------
params = pd.read_csv(PROC / "params.csv")
fit = pd.read_csv(PROC / "error_model_fit.csv").set_index("direction")
cmp_ = pd.read_csv(PROC / "model_comparison_tests.csv").set_index("direction")
sim = pd.read_csv(PROC / "type_error_sim.csv")
pc = pd.read_csv(PROC / "pc_exact.csv")
conj = pd.read_csv(PROC / "conjunctions.csv")
tle = pd.read_csv(PROC / "tle_clean.csv")
res = pd.read_csv(PROC / "residuals.csv")

# calibration numbers come from r/13 (not committed as a CSV); values from the verified run
calib = {"radial": (0.9801671, 0.9685835), "along": (0.9816107, 0.9583061), "cross": (0.9468601, 0.9498505)}

n_jumps = int(tle["jump"].astype(bool).sum()) if "jump" in tle else None


def note(text, who=None):
    """Yellow writing note: the team replaces it with their own words."""
    p = d.add_paragraph(style="Body Text")
    r = p.add_run(("✎ [" + who + "] " if who else "✎ ") + text)
    r.italic = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def bullets(items, who=None):
    for i, t in enumerate(items):
        note(("• " if i or not who else "") + t, who if i == 0 else None)


def h1(t):
    d.add_paragraph(t, style="Heading 1")


def h2(t):
    d.add_paragraph(t, style="Heading 2")


def eq(text, n):
    # the template's equation style uses the Symbol font, which garbles Unicode maths
    p = d.add_paragraph(style="equation")
    r = p.add_run("\t" + text + "\t(" + str(n) + ")")
    rpr = r._r.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), "Cambria Math")
    rpr.insert(0, fonts)
    r.font.size = Pt(9)


def body(t):
    d.add_paragraph(t, style="Body Text")


def set_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def table(title, header, rows, widths, footnote=None):
    d.add_paragraph(title, style="table head").paragraph_format.keep_with_next = True
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = 1
    set_borders(t)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.width = Inches(widths[j])
        c.paragraphs[0].style = d.styles["table col head"]
        c.paragraphs[0].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            c.width = Inches(widths[j])
            c.paragraphs[0].style = d.styles["table copy"]
            c.paragraphs[0].text = str(v)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    if footnote:
        d.add_paragraph(footnote, style="table footnote")
    else:
        d.add_paragraph(style="Normal")


def figure(path, caption, width=3.25):
    p = d.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    d.add_paragraph(caption, style="figure caption")


def pct(x):
    return f"{100 * x:.1f}%"

